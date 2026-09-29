"""XLSX ingestion preserves row context and explicitly reports omissions."""

import io
import zipfile

import pytest
from openpyxl import Workbook

from tap.modules.knowledge.adapters.document_parsers import ParserRegistry
from tap.modules.knowledge.domain.documents import DocumentParseRejected, DocumentSource, MediaType


def _bytes(book: Workbook) -> bytes:
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _parse(content: bytes):
    return ParserRegistry().parse(DocumentSource("cases.xlsx", MediaType.XLSX, content))


def test_rows_repeat_headers_and_preserve_sheet_cell_provenance():
    book = Workbook()
    sheet = book.active
    sheet.title = "Cases"
    sheet.append(["Case", "Result"])
    sheet.append(["Login", "Pass"])
    sheet.append(["Logout", "Fail"])
    result = _parse(_bytes(book))
    assert [b.text for b in result.blocks] == [
        "Case\tResult\nLogin\tPass",
        "Case\tResult\nLogout\tFail",
    ]
    assert [i.locator for i in result.parse_inventory] == [
        "worksheet:'Cases'!A2:B2;header:A1:B1",
        "worksheet:'Cases'!A3:B3;header:A1:B1",
    ]
    assert all(i.original_excerpt is None for i in result.parse_inventory)
    assert all(i.original_alignment_reason for i in result.parse_inventory)


def test_hidden_merged_and_uncached_formula_content_needs_review():
    book = Workbook()
    sheet = book.active
    sheet.append(["Case", "Value"])
    sheet.append(["Hidden", "=1+2"])
    sheet.row_dimensions[2].hidden = True
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "Merged"
    book.create_sheet("Secret").sheet_state = "hidden"
    result = _parse(_bytes(book))
    reasons = {i.reason for i in result.parse_inventory}
    assert {
        "hidden-worksheet",
        "hidden-row",
        "merged-cells",
        "formula-without-cached-value",
    } <= reasons
    assert all("=1+2" not in b.text for b in result.blocks)


def test_sparse_huge_dimensions_rejected_before_row_iteration():
    book = Workbook()
    book.active["XFD1048576"] = "far away"
    with pytest.raises(DocumentParseRejected, match="document-too-complex"):
        _parse(_bytes(book))


def test_external_links_and_macro_payloads_rejected():
    book = Workbook()
    book.active["A1"] = "link"
    book.active["A1"].hyperlink = "https://example.com"
    with pytest.raises(DocumentParseRejected, match="invalid-document"):
        _parse(_bytes(book))
    payload = io.BytesIO(_bytes(Workbook()))
    with zipfile.ZipFile(payload, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(DocumentParseRejected, match="invalid-document"):
        _parse(payload.getvalue())


def test_long_rows_keep_headers_on_every_bounded_fragment():
    book = Workbook()
    book.active.append(["Description"])
    book.active.append(["说明" * 600])
    result = _parse(_bytes(book))
    assert len(result.blocks) > 1
    assert all(b.text.startswith("A: Description\n") for b in result.blocks)
    assert all(len(b.text.encode("utf-8")) <= 512 for b in result.blocks)
    assert (
        "".join(b.text.split("\n", 1)[1].removeprefix("[continued]\n") for b in result.blocks)
        == "说明" * 600
    )


def test_archive_expansion_and_legacy_extension_are_rejected():
    payload = io.BytesIO(_bytes(Workbook()))
    with zipfile.ZipFile(payload, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/oversized.xml", b"x" * 1000000)
    with pytest.raises(DocumentParseRejected, match="document-too-complex"):
        _parse(payload.getvalue())
    with pytest.raises(DocumentParseRejected, match="unsupported-document"):
        ParserRegistry().parse(DocumentSource("legacy.xls", MediaType.XLSX, _bytes(Workbook())))


def test_cached_formulas_and_wide_headers_are_visible_but_require_review():
    book = Workbook()
    book.active.append(["H" * 300, "Result"])
    book.active.append(["Input", "=1+2"])
    book.active.column_dimensions["A"].hidden = True
    original = io.BytesIO(_bytes(book))
    output = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>")
            target.writestr(entry.filename, data)
    result = _parse(output.getvalue())
    assert {"formula-cached-value-unverified", "oversized-header-context", "hidden-column"} <= {
        item.reason for item in result.parse_inventory
    }
    assert "Input\t3" in result.blocks[0].text


def test_xml_entities_are_rejected_before_workbook_loading():
    original = io.BytesIO(_bytes(Workbook()))
    output = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                data = b'<!DOCTYPE workbook [<!ENTITY injected "unsafe">]>' + data
            target.writestr(entry.filename, data)
    with pytest.raises(DocumentParseRejected, match="invalid-document"):
        _parse(output.getvalue())


def test_long_second_cell_continuations_retain_column_and_header_identity():
    book = Workbook()
    book.active.append(["Case", "Expected result"])
    book.active.append(["Login", "Outcome " * 200])
    result = _parse(_bytes(book))
    continuations = [block for block in result.blocks if "[continued]" in block.text]
    assert continuations
    assert all(block.text.startswith("B: Expected result\n") for block in continuations)
    assert all(
        "cell:B2" in item.locator for item in result.parse_inventory if "fragment:2" in item.locator
    )


def test_formatted_values_require_review_instead_of_claiming_display_equivalence():
    book = Workbook()
    book.active.append(["Discount", "Account"])
    book.active.append([0.05, 123])
    book.active["A2"].number_format = "0%"
    book.active["B2"].number_format = "000000"
    result = _parse(_bytes(book))
    issues = [
        item for item in result.parse_inventory if item.reason == "cell-display-format-not-rendered"
    ]
    assert len(issues) == 2
    assert "0.05 [Excel format: 0%]" in result.blocks[0].text
    assert "123 [Excel format: 000000]" in result.blocks[0].text
    assert {item.locator for item in issues} == {
        "worksheet:'Sheet'!A2;format",
        "worksheet:'Sheet'!B2;format",
    }
