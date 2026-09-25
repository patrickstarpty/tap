"""Behavioral coverage for durable document completeness inventory."""

from __future__ import annotations

import base64
import io
import json
import zipfile
from xml.sax.saxutils import escape

import pytest
from docx import Document

from tap.modules.knowledge.adapters.artifact_codecs import (
    decode_normalized_artifact,
    encode_normalized_artifact,
)
from tap.modules.knowledge.adapters.document_chunker import StructuralChunker
from tap.modules.knowledge.adapters.document_parsers import ParserRegistry
from tap.modules.knowledge.domain.documents import (
    BlockKind,
    DocumentId,
    DocumentParseRejected,
    DocumentSource,
    MediaType,
    NormalizedArtifact,
    NormalizedBlock,
    RevisionId,
    canonical_sha256,
    revision_id_for,
)
from tap.modules.knowledge.domain.parse_inventory import (
    MAX_ORIGINAL_EXCERPT_BYTES,
    OriginalExcerptRange,
    ParseInventoryItem,
    ParseInventoryKind,
    ParseInventoryStatus,
    parse_inventory_digest,
    parser_config_digest,
)


def _pdf_with_text(*pages: str) -> bytes:
    return _pdf_with_streams(
        *(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1") for text in pages)
    )


def _pdf_with_streams(*streams: bytes, dictionary_prefix: bytes = b"") -> bytes:
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    page_ids = [4 + index * 2 for index in range(len(streams))]
    kids = " ".join(f"{item} 0 R" for item in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(streams)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for index, stream in enumerate(streams):
        page_id = page_ids[index]
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> "
                f"/MediaBox [0 0 612 792] /Contents {page_id + 1} 0 R >>"
            ).encode()
        )
        objects.append(
            b"<< "
            + dictionary_prefix
            + f"/Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream"
        )
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(body))
        body.extend(f"{number} 0 obj\n".encode())
        body.extend(value)
        body.extend(b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    body.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    body.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(body)


def _pdf_with_inline_image_decoy_and_form() -> bytes:
    visible = b"BT /F1 12 Tf 72 720 Td (Visible clause) Tj ET"
    page_stream = b"q BI /W 1 /H 1 /CS /G /BPC 8 ID " + visible + b" EI Q /Fm1 Do"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R] /Count 1 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            b"<< /Type /Page /Parent 2 0 R /Resources "
            b"<< /Font << /F1 3 0 R >> /XObject << /Fm1 6 0 R >> >> "
            b"/MediaBox [0 0 612 792] /Contents 5 0 R >>"
        ),
        f"<< /Length {len(page_stream)} >>\nstream\n".encode() + page_stream + b"\nendstream",
        (
            b"<< /Type /XObject /Subtype /Form /BBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> "
            + f"/Length {len(visible)} >>\nstream\n".encode()
            + visible
            + b"\nendstream"
        ),
    ]
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(body))
        body.extend(f"{number} 0 obj\n".encode())
        body.extend(value)
        body.extend(b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    body.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    body.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(body)


def _identified_source(filename: str, media_type: MediaType, content: bytes) -> DocumentSource:
    document_id = DocumentId("doc_inventory")
    return DocumentSource(
        filename,
        media_type,
        content,
        document_id,
        revision_id_for(document_id, canonical_sha256(content), "tapper-parser-v1"),
    )


def _stored_docx(*paragraphs: tuple[str, str]) -> bytes:
    body = "".join(
        '<w:p><w:pPr><w:pStyle w:val="%s"/></w:pPr><w:r><w:t>%s</w:t></w:r></w:p>'
        % (style, escape(text))
        for style, text in paragraphs
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}<w:sectPr/></w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.'
        'wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        'officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)
    return payload.getvalue()


def _stored_docx_body(body: str) -> bytes:
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}<w:sectPr/></w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.'
        'wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        'officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)
    return payload.getvalue()


def test_mixed_pdf_inventory_does_not_hide_a_page_without_extractable_text() -> None:
    """Skipping a blank/scanned page would falsely mark a mixed PDF complete."""
    artifact = ParserRegistry().parse(
        _identified_source("mixed.pdf", MediaType.PDF, _pdf_with_text("Approved rule", ""))
    )

    pages = [item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PAGE]
    assert [(item.locator, item.status, item.reason) for item in pages] == [
        ("page:1", ParseInventoryStatus.PARSED, None),
        ("page:2", ParseInventoryStatus.FAILED, "ocr-required"),
    ]
    assert {block.inventory_item_id for block in artifact.blocks} == {pages[0].item_id}
    assert artifact.parser_config_digest == parser_config_digest(MediaType.PDF.value)
    assert artifact.parse_inventory_digest == parse_inventory_digest(artifact.parse_inventory)


def test_text_inventory_persists_a_bounded_unicode_byte_range_for_a_late_item() -> None:
    """A code-point offset or document-head fallback would select the wrong original bytes."""
    prefix = ("前置🙂" * 180_000).encode()
    target = "后部精确条款🙂"
    content = prefix + b"\n\n" + target.encode()

    artifact = ParserRegistry().parse(_identified_source("late.txt", MediaType.TEXT, content))

    item = artifact.parse_inventory[-1]
    assert item.original_excerpt is not None
    span = item.original_excerpt
    assert content[span.start_byte : span.end_byte].decode() == target
    assert span.end_byte - span.start_byte <= MAX_ORIGINAL_EXCERPT_BYTES
    assert span.source_digest == canonical_sha256(content)
    assert span.excerpt_digest == canonical_sha256(target.encode())


def test_text_provenance_keeps_repeated_nfc_equivalent_items_on_their_source_nodes() -> None:
    """A normalized whole-file search can bind decomposed text to a different item."""
    decomposed = "e\N{COMBINING ACUTE ACCENT}"
    content = f"{decomposed}\n\né\n\n{decomposed}".encode()

    artifact = ParserRegistry().parse(_identified_source("unicode.txt", MediaType.TEXT, content))

    excerpts = [
        content[item.original_excerpt.start_byte : item.original_excerpt.end_byte]
        for item in artifact.parse_inventory
        if item.original_excerpt is not None
    ]
    assert excerpts == [decomposed.encode(), "é".encode(), decomposed.encode()]
    assert [block.text for block in artifact.blocks] == ["é", "é", "é"]


def test_text_provenance_is_unsupported_when_normalization_has_no_safe_boundary_map() -> None:
    """Normalization across independent Hangul Jamo cannot be guessed as raw byte bounds."""
    content = "\N{HANGUL CHOSEONG KIYEOK}\N{HANGUL JUNGSEONG A}".encode()

    artifact = ParserRegistry().parse(_identified_source("jamo.txt", MediaType.TEXT, content))

    item = artifact.parse_inventory[0]
    assert artifact.blocks[0].text == "가"
    assert item.original_excerpt is None
    assert item.original_alignment_reason == "source-text-not-stably-addressable"


def test_text_provenance_does_not_scan_the_source_once_per_inventory_item() -> None:
    """A parser with many items must not regress to items multiplied by source scans."""

    class CountingBytes(bytes):
        find_calls = 0

        def find(self, *args, **kwargs):  # type: ignore[no-untyped-def]
            type(self).find_calls += 1
            if type(self).find_calls > 2:
                raise AssertionError("parser repeatedly scanned the complete source")
            return super().find(*args, **kwargs)

    content = CountingBytes(
        "\n\n".join(f"item-{index}-" + "x" * 6000 for index in range(200)).encode()
    )

    artifact = ParserRegistry().parse(_identified_source("many.txt", MediaType.TEXT, content))

    assert len(artifact.parse_inventory) == 200
    assert CountingBytes.find_calls == 0


def test_text_provenance_hashes_a_large_source_once_per_parse(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Excerpt construction must reuse the artifact digest instead of rehashing per item."""
    from tap.modules.knowledge.adapters import document_parsers

    content = "\n\n".join(f"item-{index}-" + "x" * 6000 for index in range(200)).encode()
    source = _identified_source("many.txt", MediaType.TEXT, content)
    original_sha256 = document_parsers.canonical_sha256
    source_hash_calls = 0

    def counted_sha256(value: bytes) -> str:
        nonlocal source_hash_calls
        if value is source.content:
            source_hash_calls += 1
        return original_sha256(value)

    monkeypatch.setattr(document_parsers, "canonical_sha256", counted_sha256)

    artifact = ParserRegistry().parse(source)

    assert len(artifact.parse_inventory) == 200
    assert len(source.content) > 1_000_000
    assert source_hash_calls == 1


def test_markdown_original_range_keeps_source_markup_instead_of_relabeling_extracted_text() -> None:
    """The normalized heading title must never masquerade as the Markdown original."""
    content = "# 原始标题🙂\r\n\r\n正文".encode()

    artifact = ParserRegistry().parse(_identified_source("guide.md", MediaType.MARKDOWN, content))

    heading = artifact.parse_inventory[0]
    assert heading.original_excerpt is not None
    span = heading.original_excerpt
    assert content[span.start_byte : span.end_byte].decode() == "# 原始标题🙂"
    assert artifact.blocks[0].text == "原始标题🙂"


def test_markdown_provenance_keeps_repeated_nfc_equivalent_nodes_distinct() -> None:
    """Markdown syntax positions, not normalized value equality, choose the original item."""
    decomposed = "e\N{COMBINING ACUTE ACCENT}"
    content = f"# {decomposed}\n\n# é\n\n# {decomposed}".encode()

    artifact = ParserRegistry().parse(_identified_source("unicode.md", MediaType.MARKDOWN, content))

    excerpts = [
        content[item.original_excerpt.start_byte : item.original_excerpt.end_byte]
        for item in artifact.parse_inventory
        if item.original_excerpt is not None
    ]
    assert excerpts == [
        f"# {decomposed}".encode(),
        "# é".encode(),
        f"# {decomposed}".encode(),
    ]
    assert [block.text for block in artifact.blocks] == ["é", "é", "é"]


def test_pdf_and_docx_only_claim_original_alignment_for_an_exact_source_span() -> None:
    """Container text that cannot be tied to exact immutable bytes must stay unsupported."""
    pdf = ParserRegistry().parse(
        _identified_source("policy.pdf", MediaType.PDF, _pdf_with_text("Exact PDF clause"))
    )
    pdf_item = next(
        item for item in pdf.parse_inventory if item.status is ParseInventoryStatus.PARSED
    )
    assert pdf_item.original_excerpt is not None

    docx = ParserRegistry().parse(
        _identified_source(
            "policy.docx",
            MediaType.DOCX,
            _stored_docx(("Normal", "Exact DOCX clause"), ("Normal", "A & B")),
        )
    )
    exact, escaped = [
        item for item in docx.parse_inventory if item.kind is ParseInventoryKind.PARAGRAPH
    ]
    assert exact.original_excerpt is not None
    assert escaped.original_excerpt is not None
    escaped_span = escaped.original_excerpt
    assert (
        _stored_docx(("Normal", "Exact DOCX clause"), ("Normal", "A & B"))[
            escaped_span.start_byte : escaped_span.end_byte
        ]
        == b"A &amp; B"
    )


def test_pdf_provenance_uses_the_text_showing_token_not_an_identical_comment() -> None:
    """A PDF comment must not defeat or impersonate the displayed TJ token."""
    payload = _pdf_with_streams(
        b"% Exact PDF clause\nBT /F1 12 Tf 72 720 Td [(Exact PDF clause)] TJ ET"
    )

    artifact = ParserRegistry().parse(_identified_source("operator.pdf", MediaType.PDF, payload))

    item = next(item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PAGE)
    assert item.original_excerpt is not None
    span = item.original_excerpt
    assert payload[span.start_byte : span.end_byte] == b"Exact PDF clause"
    assert payload[: span.start_byte].endswith(b"[(")


def test_pdf_provenance_skips_stream_decoy_inside_object_dictionary() -> None:
    """Only the stream keyword after the parsed object dictionary can anchor text."""
    stream = b"BT /F1 12 Tf 72 720 Td (Visible clause) Tj ET"
    payload = _pdf_with_streams(
        stream,
        dictionary_prefix=b"/Decoy (stream\n" + stream + b") ",
    )

    artifact = ParserRegistry().parse(_identified_source("decoy.pdf", MediaType.PDF, payload))

    item = next(item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PAGE)
    assert item.original_excerpt is not None
    span = item.original_excerpt
    actual_stream_start = payload.rfind(b"\nstream\n") + len(b"\nstream\n")
    assert span.start_byte >= actual_stream_start
    assert payload[span.start_byte : span.end_byte] == b"Visible clause"


def test_pdf_provenance_is_unsupported_when_inline_image_payload_contains_fake_text() -> None:
    """Inline image bytes must never impersonate text rendered by a Form XObject."""
    payload = _pdf_with_inline_image_decoy_and_form()

    artifact = ParserRegistry().parse(_identified_source("inline.pdf", MediaType.PDF, payload))

    item = next(item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PAGE)
    assert artifact.blocks[0].text == "Visible clause"
    assert item.original_excerpt is None
    assert item.original_alignment_reason == "source-text-not-stably-addressable"


def test_pdf_provenance_lexer_stops_before_inline_image_payload() -> None:
    """Fail-closed inline-image detection must not inspect attacker-controlled image bytes."""
    from tap.modules.knowledge.adapters import document_parsers

    class PayloadGuard(bytes):
        def __getitem__(self, key):  # type: ignore[no-untyped-def]
            if isinstance(key, int) and key > 2:
                raise AssertionError("inline image payload was scanned")
            return super().__getitem__(key)

    content = PayloadGuard(b"BI attacker-controlled payload with (Visible clause) Tj")

    assert document_parsers._pdf_tokens(content) is None


def test_pdf_provenance_remains_available_without_inline_image() -> None:
    """The inline-image guard must not disable a direct text-showing page stream."""
    payload = _pdf_with_text("Visible clause")

    artifact = ParserRegistry().parse(_identified_source("plain.pdf", MediaType.PDF, payload))

    item = next(item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PAGE)
    assert item.original_excerpt is not None
    span = item.original_excerpt
    assert payload[span.start_byte : span.end_byte] == b"Visible clause"


def test_docx_split_runs_never_bind_to_matching_hidden_attribute_text() -> None:
    """A bookmark/attribute match is not provenance for visible split-run text."""
    payload = _stored_docx_body(
        '<w:p><w:bookmarkStart w:id="0" w:name="Split clause"/>'
        "<w:r><w:t>Split </w:t></w:r><w:r><w:t>clause</w:t></w:r>"
        '<w:bookmarkEnd w:id="0"/></w:p>'
    )

    artifact = ParserRegistry().parse(_identified_source("split.docx", MediaType.DOCX, payload))

    item = next(
        item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PARAGRAPH
    )
    assert item.original_excerpt is None
    assert item.original_alignment_reason == "source-text-not-stably-addressable"


def test_docx_text_tag_end_ignores_greater_than_inside_quoted_attribute() -> None:
    """A quoted greater-than character must not become the visible text start."""
    payload = _stored_docx_body('<w:p><w:r><w:t data-x=">">Visible clause</w:t></w:r></w:p>')

    artifact = ParserRegistry().parse(_identified_source("quoted.docx", MediaType.DOCX, payload))

    item = next(
        item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.PARAGRAPH
    )
    assert item.original_excerpt is not None
    span = item.original_excerpt
    assert payload[span.start_byte : span.end_byte] == b"Visible clause"


def test_original_excerpt_range_rejects_malformed_or_unbounded_offsets() -> None:
    """A malformed persisted range must fail closed before any object-store read."""
    with pytest.raises(ValueError, match="range"):
        OriginalExcerptRange(
            source_digest=canonical_sha256(b"text"),
            start_byte=4,
            end_byte=3,
            excerpt_digest=canonical_sha256(b""),
        )
    with pytest.raises(ValueError, match="bound"):
        OriginalExcerptRange(
            source_digest=canonical_sha256(b"x" * (MAX_ORIGINAL_EXCERPT_BYTES + 1)),
            start_byte=0,
            end_byte=MAX_ORIGINAL_EXCERPT_BYTES + 1,
            excerpt_digest=canonical_sha256(b"x" * (MAX_ORIGINAL_EXCERPT_BYTES + 1)),
        )


def test_docx_inventory_records_table_and_embedded_image_separately() -> None:
    """Flattening a table or silently dropping an image destroys completeness evidence."""
    document = Document()
    document.add_paragraph("Approved rule")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Count"
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document.add_picture(io.BytesIO(png))
    payload = io.BytesIO()
    document.save(payload)

    artifact = ParserRegistry().parse(
        _identified_source("mixed.docx", MediaType.DOCX, payload.getvalue())
    )

    table_item = next(
        item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.TABLE
    )
    image_item = next(
        item for item in artifact.parse_inventory if item.kind is ParseInventoryKind.IMAGE
    )
    assert (table_item.locator, table_item.status) == (
        "table:1",
        ParseInventoryStatus.PARSED,
    )
    assert (image_item.status, image_item.reason) == (
        ParseInventoryStatus.NEEDS_REVIEW,
        "image-not-text-extracted",
    )
    table_block = next(block for block in artifact.blocks if block.kind is BlockKind.TABLE_TEXT)
    assert table_block.inventory_item_id == table_item.item_id


def test_invalid_text_encoding_returns_a_failed_inventory_item() -> None:
    """Rejecting bytes without a durable item would make the completeness gap disappear."""
    artifact = ParserRegistry().parse(
        _identified_source("broken.txt", MediaType.TEXT, b"valid-prefix\xff")
    )

    assert artifact.blocks == ()
    assert [(item.status, item.reason, item.locator) for item in artifact.parse_inventory] == [
        (ParseInventoryStatus.FAILED, "invalid-encoding", "document:text")
    ]
    with pytest.raises(DocumentParseRejected, match="^invalid-document$"):
        StructuralChunker().chunk(artifact)


def test_image_only_pdf_keeps_the_ocr_remediation_at_the_chunking_boundary() -> None:
    """Persisting inventory must not degrade the existing safe OCR error shown to users."""
    artifact = ParserRegistry().parse(
        _identified_source("scan.pdf", MediaType.PDF, _pdf_with_text(""))
    )

    with pytest.raises(DocumentParseRejected, match="^ocr-required$"):
        StructuralChunker().chunk(artifact)


def test_chunk_anchor_binds_to_the_inventory_item_that_produced_its_text() -> None:
    """A chunk without an inventory identity cannot prove which checked object it came from."""
    item = ParseInventoryItem.create(
        source_revision_id="rev_inventory",
        kind=ParseInventoryKind.PARAGRAPH,
        locator="paragraph:1",
        status=ParseInventoryStatus.PARSED,
        artifact_digest=canonical_sha256(b"Approved rule"),
    )
    artifact = NormalizedArtifact(
        filename="policy.txt",
        media_type=MediaType.TEXT,
        source_hash=canonical_sha256(b"Approved rule"),
        blocks=(
            NormalizedBlock(
                block_id="b_000000",
                kind=BlockKind.PARAGRAPH,
                text="Approved rule",
                heading_path=(),
                page=None,
                paragraph_index=0,
                start_offset=0,
                end_offset=13,
                inventory_item_id=item.item_id,
            ),
        ),
        document_id=DocumentId("doc_inventory"),
        revision_id=RevisionId("rev_inventory"),
        parse_inventory=(item,),
        parser_config_digest=parser_config_digest(MediaType.TEXT.value),
        parse_inventory_digest=parse_inventory_digest((item,)),
    )

    chunk = StructuralChunker().chunk(artifact)[0]

    assert f'"inventoryItemId":"{item.item_id}"' in chunk.anchor_json


def test_excluded_inventory_item_requires_both_reason_and_accountable_actor() -> None:
    """An anonymous exclusion would let missing source content evade independent review."""
    with pytest.raises(ValueError, match="actor"):
        ParseInventoryItem.create(
            source_revision_id="rev_inventory",
            kind=ParseInventoryKind.IMAGE,
            locator="image:1",
            status=ParseInventoryStatus.EXCLUDED,
            reason="decorative",
            artifact_digest=canonical_sha256(b"image"),
        )

    excluded = ParseInventoryItem.create(
        source_revision_id="rev_inventory",
        kind=ParseInventoryKind.IMAGE,
        locator="image:1",
        status=ParseInventoryStatus.EXCLUDED,
        reason="decorative",
        artifact_digest=canonical_sha256(b"image"),
        decision_actor_id="actor-reviewer",
    )

    assert excluded.status is ParseInventoryStatus.EXCLUDED
    assert excluded.decision_actor_id == "actor-reviewer"


def test_normalized_artifact_codec_preserves_inventory_and_reads_legacy_as_unreviewed() -> None:
    """An upgrade must retain new completeness facts without calling old data approved."""
    source = _identified_source("policy.txt", MediaType.TEXT, b"Approved rule")
    current = ParserRegistry().parse(source)
    assert current.revision_id is not None

    restored = decode_normalized_artifact(
        encode_normalized_artifact(str(current.revision_id), current),
        expected_revision=str(current.revision_id),
    )

    assert restored.parse_inventory == current.parse_inventory
    assert restored.parse_inventory_digest == current.parse_inventory_digest

    source_hash = canonical_sha256(b"Legacy rule")
    document_id = DocumentId("doc_legacy")
    revision_id = revision_id_for(document_id, source_hash, "tapper-parser-v1")
    payload = {
        "blocks": [
            {
                "blockId": "b_000000",
                "endOffset": 11,
                "headingPath": [],
                "kind": "paragraph",
                "page": None,
                "paragraphIndex": 0,
                "startOffset": 0,
                "text": "Legacy rule",
            }
        ],
        "documentId": str(document_id),
        "filename": "legacy.txt",
        "mediaType": "text/plain",
        "normalizedSchema": "normalized-artifact-v1",
    }
    canonical_payload = (
        json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
        + b"\n"
    )
    envelope = {
        "blockCount": 1,
        "payload": payload,
        "payloadSha256": canonical_sha256(canonical_payload),
        "revisionId": str(revision_id),
        "schemaVersion": "normalized-v1",
        "sourceContentHash": source_hash,
    }
    legacy_bytes = (
        json.dumps(envelope, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
        + b"\n"
    )

    legacy = decode_normalized_artifact(legacy_bytes, expected_revision=str(revision_id))

    assert legacy.blocks[0].text == "Legacy rule"
    assert legacy.parse_inventory[0].status is ParseInventoryStatus.NEEDS_REVIEW
    assert legacy.parse_inventory[0].reason == "historical-unreviewed"
    assert legacy.parse_inventory[0].original_excerpt is None
    assert legacy.parse_inventory[0].original_alignment_reason is None
